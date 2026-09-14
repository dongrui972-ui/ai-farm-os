from __future__ import annotations

from pathlib import Path
from typing import Iterable, Sequence

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "docs" / "一级芯界_AI_Farm_OS_系统文档.docx"
HERO = ROOT / "frontend" / "assets" / "img" / "farm-real" / "hero-xinjiang.jpg"

BLACK = "111111"
GREEN = "1E5A3B"
GREEN_DARK = "15432C"
GREEN_LIGHT = "EAF3ED"
GREEN_PALE = "F5F9F6"
GRAY = "5D665F"
GRAY_LIGHT = "F2F4F2"
GRAY_BORDER = "D9D9D9"
AMBER = "9B5B00"
RED = "9B2C2C"


def set_east_asia_font(run, name: str = "Microsoft YaHei") -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)


def set_language(element, value: str = "zh-CN") -> None:
    rpr = element.get_or_add_rPr()
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rpr.append(lang)
    lang.set(qn("w:val"), value)
    lang.set(qn("w:eastAsia"), value)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color: str = GRAY_BORDER, size: str = "6") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = tr_pr.find(qn("w:cantSplit"))
    if cant_split is None:
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)


def set_cell_width(cell, width_inches: float) -> None:
    cell.width = Inches(width_inches)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_text(cell, text: str, *, bold: bool = False, color: str = BLACK, size: float = 8.7) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(str(text))
    set_east_asia_font(run)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def add_table(
    doc: Document,
    headers: Sequence[str],
    rows: Iterable[Sequence[str]],
    widths: Sequence[float] | None = None,
    *,
    font_size: float = 8.6,
) :
    rows = list(rows)
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header = table.rows[0]
    set_repeat_table_header(header)
    prevent_row_split(header)
    for i, text in enumerate(headers):
        set_cell_text(header.cells[i], text, bold=True, color="FFFFFF", size=font_size)
        set_cell_shading(header.cells[i], GREEN_DARK)
        if widths:
            set_cell_width(header.cells[i], widths[i])
    for row_index, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for i, value in enumerate(values):
            set_cell_text(row.cells[i], str(value), size=font_size)
            if row_index % 2 == 1:
                set_cell_shading(row.cells[i], GREEN_PALE)
            if widths:
                set_cell_width(row.cells[i], widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_east_asia_font(run)
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(GRAY)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    suffix = paragraph.add_run(" 页")
    set_east_asia_font(suffix)
    suffix.font.size = Pt(8)
    suffix.font.color.rgb = RGBColor.from_string(GRAY)


def set_picture_alt(inline_shape, alt_text: str) -> None:
    nodes = inline_shape._inline.xpath(".//pic:cNvPr")
    if nodes:
        nodes[0].set("descr", alt_text)
        nodes[0].set("title", "新疆智慧农田资料图")


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="Caption")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(text)


def add_h1(doc: Document, text: str, *, new_page: bool = False) -> None:
    if new_page:
        doc.add_page_break()
    p = doc.add_paragraph(style="Heading 1")
    p.add_run(text)


def add_h2(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="Heading 2")
    p.add_run(text)


def add_h3(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="Heading 3")
    p.add_run(text)


def add_para(doc: Document, text: str, *, lead: str | None = None, italic: bool = False) -> None:
    p = doc.add_paragraph()
    if lead and text.startswith(lead):
        r1 = p.add_run(lead)
        r1.bold = True
        set_east_asia_font(r1)
        r2 = p.add_run(text[len(lead):])
        set_east_asia_font(r2)
        r2.italic = italic
    else:
        r = p.add_run(text)
        set_east_asia_font(r)
        r.italic = italic


def add_bullets(doc: Document, items: Iterable[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.24)
        p.paragraph_format.first_line_indent = Inches(-0.16)
        r = p.add_run(item)
        set_east_asia_font(r)


def add_numbered(doc: Document, items: Iterable[str]) -> None:
    numbering = doc.part.numbering_part.element
    abstract_id = None
    for abstract in numbering.findall(qn("w:abstractNum")):
        for level in abstract.findall(qn("w:lvl")):
            ilvl_value = level.get(qn("w:ilvl"))
            num_fmt = level.find(qn("w:numFmt"))
            if ilvl_value == "0" and num_fmt is not None and num_fmt.get(qn("w:val")) == "decimal":
                abstract_id = abstract.get(qn("w:abstractNumId"))
                break
        if abstract_id is not None:
            break
    if abstract_id is None:
        abstract_id = "0"
    current_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    num_id = max(current_ids, default=0) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    level_override = OxmlElement("w:lvlOverride")
    level_override.set(qn("w:ilvl"), "0")
    start_override = OxmlElement("w:startOverride")
    start_override.set(qn("w:val"), "1")
    level_override.append(start_override)
    num.append(level_override)
    numbering.append(num)
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.left_indent = Inches(0.26)
        p.paragraph_format.first_line_indent = Inches(-0.18)
        p_pr = p._p.get_or_add_pPr()
        num_pr = p_pr.find(qn("w:numPr"))
        if num_pr is None:
            num_pr = OxmlElement("w:numPr")
            p_pr.append(num_pr)
        ilvl_node = OxmlElement("w:ilvl")
        ilvl_node.set(qn("w:val"), "0")
        num_id_node = OxmlElement("w:numId")
        num_id_node.set(qn("w:val"), str(num_id))
        num_pr.extend([ilvl_node, num_id_node])
        r = p.add_run(item)
        set_east_asia_font(r)


def add_code(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="Code")
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(5)
    run = p.add_run(text)
    run.font.name = "Cascadia Mono"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(8.3)
    run.font.color.rgb = RGBColor.from_string(BLACK)


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.2)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.line_spacing = 1.22
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.widow_control = True

    title = styles["Title"]
    title.font.name = "Microsoft YaHei"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    title.font.size = Pt(26)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(BLACK)
    title.paragraph_format.space_after = Pt(8)
    title.paragraph_format.keep_with_next = True
    title_ppr = title._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    subtitle = styles["Subtitle"]
    subtitle.font.name = "Microsoft YaHei"
    subtitle._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    subtitle.font.size = Pt(13)
    subtitle.font.color.rgb = RGBColor.from_string(GRAY)
    subtitle.font.italic = False
    subtitle.paragraph_format.space_after = Pt(18)
    subtitle.paragraph_format.keep_with_next = True

    for name, size, before, after in (
        ("Heading 1", 18, 14, 9),
        ("Heading 2", 13, 10, 5),
        ("Heading 3", 11, 7, 3),
    ):
        style = styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    caption = styles["Caption"]
    caption.font.name = "Microsoft YaHei"
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    caption.font.size = Pt(8.2)
    caption.font.color.rgb = RGBColor.from_string(GRAY)
    caption.font.italic = False
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(7)

    try:
        code = styles["Code"]
    except KeyError:
        code = styles.add_style("Code", WD_STYLE_TYPE.PARAGRAPH)
    code.font.name = "Cascadia Mono"
    code._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    code.font.size = Pt(8.3)
    code.paragraph_format.line_spacing = 1.0
    code.paragraph_format.widow_control = False

    for style_name in ("List Bullet", "List Number"):
        st = styles[style_name]
        st.font.name = "Microsoft YaHei"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        st.font.size = Pt(9.8)
        st.paragraph_format.space_after = Pt(2)


def build_document() -> None:
    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.66)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.28)
    section.different_first_page_header_footer = True

    header = section.header
    hp = header.paragraphs[0]
    hp.text = "一级芯界 AI Farm OS · 系统文档"
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in hp.runs:
        set_east_asia_font(run)
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(GRAY)
    add_page_number(section.footer.paragraphs[0])

    props = doc.core_properties
    props.title = "一级芯界 AI Farm OS 系统文档"
    props.subject = "系统设计、功能、数据、安全、部署与验收说明"
    props.author = "一级芯界 AI Farm OS 项目组"
    props.keywords = "智慧农业, AI Farm OS, 农业决策, 农机安全, 系统设计"
    props.comments = "依据当前代码仓库与 FINAL_SPEC v1.2 编制"

    # 封面
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.add_run("一级芯界 AI Farm OS")
    sp = doc.add_paragraph(style="Subtitle")
    sp.add_run("系统设计 运维与验收说明书")
    if HERO.exists():
        pic_p = doc.add_paragraph()
        pic_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        shape = pic_p.add_run().add_picture(str(HERO), width=Inches(6.9))
        set_picture_alt(shape, "新疆智慧农田全景资料图，展示成片农田、道路与远处山脉，用于说明系统应用场景")
        add_caption(doc, "新疆智慧农田资料图　本地静态素材，不代表实时画面或设备回传")

    add_table(
        doc,
        ["文档属性", "内容"],
        [
            ("文档版本", "1.0"),
            ("系统基线", "FINAL_SPEC Final v1.2；前端资源 v155"),
            ("编制日期", "2026-09-14"),
            ("适用范围", "本地决策支持、交互验证与受控作业沙箱"),
            ("适用读者", "项目负责人、农业专家、产品、研发、测试、安全、运维与设备接入人员"),
        ],
        [1.45, 5.35],
        font_size=9.0,
    )
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    r = p.add_run("重要边界：本系统当前不直接控制真实泵站、无人机、农机或机器人，也不把仿真数据当作生产遥测。")
    set_east_asia_font(r)
    r.font.size = Pt(9.2)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(RED)

    # 目录与摘要
    add_h1(doc, "文档使用说明", new_page=True)
    add_para(doc, "本说明书依据当前代码、产品终版规范、专家联审规则和自动化测试结果编制。它既描述已经可以运行的功能，也标记尚未生产化的目标能力。任何涉及真实灌溉、施药、农机运动或仓储执行的功能，都必须先满足本文件列出的证据、权限和物理安全条件。")
    add_h2(doc, "执行摘要")
    add_bullets(doc, [
        "系统已形成可运行的中文农业决策界面，提供农场主、专家、监管三类视图，覆盖看田、任务、水肥、全季装备、收贮、智能研判、数据资产与系统架构。",
        "当前主业务以浏览器本地规则引擎和固定仿真数据运行；可选 FastAPI + SQLite 后端对外声明 business_api=partial，因而不会被前端误判为完整生产数据源。",
        "写操作默认锁定，关键证据缺失时返回 NO_GO；角色切换仅改变呈现，不等同于服务器授权。",
        "当前版本适合产品评审、流程验证、交互验收、接口联调和安全规则讨论，不适合无人值守生产控制、法定计量、结算或监管取证。",
        "发布门禁已在 2026-09-14 复核：JavaScript 语法、Python 编译、PowerShell 解析与 34 项自动化测试全部通过；存在 1 条第三方弃用告警。",
    ])
    add_h2(doc, "章节导航")
    contents = [
        ("1", "系统定位与范围"), ("2", "用户角色与信息架构"),
        ("3", "总体架构"), ("4", "核心业务闭环"),
        ("5", "功能模块"), ("6", "首页与态势大屏"),
        ("7", "数据架构与质量合同"), ("8", "接口与集成"),
        ("9", "安全与权限"), ("10", "农业决策与农机安全"),
        ("11", "部署运行"), ("12", "测试与验收"),
        ("13", "生产化路线"), ("14", "运维与故障处理"),
        ("A", "接口清单"), ("B", "数据表与代码索引"),
    ]
    add_table(doc, ["序号", "主题", "序号", "主题"], [
        (contents[i][0], contents[i][1], contents[i + 8][0], contents[i + 8][1])
        for i in range(8)
    ], [0.55, 2.85, 0.55, 2.85], font_size=8.6)

    # 1
    add_h1(doc, "1 系统定位与范围", new_page=True)
    add_h2(doc, "1.1 产品定位")
    add_para(doc, "一级芯界 AI Farm OS 是面向规模化农场的农业决策支持与受控作业沙箱。产品主线是“看田—形成建议—补齐证据—人工把关—仿真派工—记录结果—复盘改进”，把农艺、农机、传感、任务和审计放在同一操作闭环中。")
    add_para(doc, "当前场景基线为新疆库尔勒试验基地，面积 2000 亩，作物覆盖棉花、玉米和小麦。界面中的节水约 20%、病虫识别约 89.5%、人工巡田减少不低于 50%、亩均增收 400—600 元等数值均为目标指标，不是已达成的生产绩效。")
    add_h2(doc, "1.2 能力边界")
    add_table(doc, ["能力层级", "当前状态", "可以做什么", "不可据此声称"], [
        ("交互与流程", "已实现", "角色导航、任务优先首页、地图态势、证据对话框、仿真操作、结果留痕", "真实作业已执行或已验收"),
        ("农业规则", "已实现基础门禁", "缺证据时 NO_GO，水肥、风险和设备状态采用规则化判断", "经过田间标定的处方模型"),
        ("业务接口", "部分实现", "FastAPI + SQLite 查询、受控写入、幂等与审计", "完整业务后端或多租户 SaaS"),
        ("设备接入", "沙箱能力", "登记设备、创建仿真控制请求、机器人任务排队", "已连接真实设备、物理急停或安全 PLC"),
        ("AI 能力", "规则与可选问答", "关键词视觉判断、产量公式、风险仿真、可选 OpenAI 问答代理", "经认证的病虫诊断或无人决策"),
        ("生产目标", "未交付", "云边端、时空数据库、向量与图谱、设备签名、双人复核、不可变审计", "现版本已经达到生产等级"),
    ], [1.0, 1.0, 2.8, 2.0], font_size=8.1)
    add_h2(doc, "1.3 设计原则")
    add_numbered(doc, [
        "任务优先：首页先展示今天必须处理的事项，再提供态势和分析背景。",
        "决策可解释：每条建议必须给出原因、风险、负责人、期限和证据引用。",
        "默认拒绝：数据缺失、时效过期、质量不明或空间坐标不合格时，不生成可执行指令。",
        "观察与命令分离：遥测、建议、审批、仿真命令、设备确认和作业验收分别记录。",
        "人机协同：高风险动作必须由有资质人员复核，远程停止不能替代现场物理急停。",
        "一份事实源：大屏、首页和详情页共享同一状态，不在展示层制造第二套业务结论。",
    ])

    # 2
    add_h1(doc, "2 用户角色与信息架构")
    add_h2(doc, "2.1 角色职责")
    add_table(doc, ["角色", "首要目标", "核心路径", "权限边界"], [
        ("农场主", "安排本周农事并跟踪完成", "看天 → 本周农事 → 采纳、改期或跳过 → 回放或停止 → 留痕", "角色切换只改变界面；生产写入仍需服务器身份与农场范围校验"),
        ("农业专家", "联合研判、补证并审核处方", "研判 → 补证 → 初审 → 人工门禁 → 仿真复盘", "审核者身份、资质范围和证据必须由服务端核验"),
        ("监管人员", "查看异常、效果和合规证据", "异常 → 指标 → 取证 → 关注", "默认只读；不能因可见按钮而获得作业权限"),
    ], [0.8, 1.65, 2.7, 1.65], font_size=8.3)
    add_h2(doc, "2.2 导航结构")
    add_table(doc, ["角色", "导航分组", "主要页面"], [
        ("农场主", "计划", "本周农事"),
        ("农场主", "田块", "看田地图、田间设备"),
        ("农场主", "作业", "浇水施肥、机具作业、播收到仓"),
        ("农场主", "账本与问答", "历年收成、问一问"),
        ("农业专家", "今日研判", "专家驾驶舱、数字孪生、历年档案"),
        ("农业专家", "审核与协同", "任务审核台、联合诊断、协作工作台、多 Agent 协同、Agent 名册、AI 分析中心"),
        ("农业专家", "处方与装备", "水肥研判、植株管控、多机联作、机器人沙箱、全季装备、设备网、厂家接入"),
        ("农业专家", "收贮与底座", "仓储初加工、数据资产、云边端架构"),
        ("监管人员", "先看异常", "监管驾驶舱、高优任务监察、田间一张图、历年档案"),
        ("监管人员", "效果与合规", "水肥节量、全季进度、收贮溯源、数据资产、政策问答"),
        ("监管人员", "运行监察", "机队作业、设备核验、协同事件、异常诊断、系统架构"),
    ], [0.85, 1.55, 4.4], font_size=8.2)
    add_h2(doc, "2.3 页面交互约定")
    add_bullets(doc, [
        "全局顶栏提供农场选择、搜索、运行模式、天气、时钟、刷新、主题、态势大屏、通知与角色入口。",
        "侧栏按业务阶段分组；当前页面和折叠状态必须清晰，菜单上方不得留下无信息的空白区域。",
        "所有可执行动作先展示风险和证据；主操作、次操作与破坏性操作在视觉和键盘顺序上保持一致。",
        "对话框支持标题、可访问名称、焦点陷阱、Esc 关闭和关闭后的焦点恢复。",
    ])

    # 3
    add_h1(doc, "3 总体架构")
    add_h2(doc, "3.1 当前运行架构")
    add_para(doc, "系统当前是静态单页前端与可选后端沙箱的组合。前端通过 FarmRuntime 探测 /api/health；只有后端同时报告 status=running 且 business_api=complete 时，业务页面才使用后端数据。当前后端诚实报告 business_api=partial，因此业务页面默认使用本地 FarmEngine，避免半成品接口被误当成生产事实。AI 问答可单独调用 /api/ai/chat。")
    add_table(doc, ["入口", "前端与路由", "规则与接口", "数据存储", "输出"], [
        ("浏览器\n127.0.0.1", "HTML + CSS + 原生 JavaScript\n三角色界面", "FarmRuntime\n健康探测与降级", "FarmEngine 固定仿真状态", "页面、对话框、CSV"),
        ("start-web.ps1", "本地静态服务\nHost 与路径白名单", "可选 AI 问答代理\n16 KiB 与限流", "环境变量中的 AI 配置", "结构化问答"),
        ("FastAPI", "REST API\n业务能力 partial", "写入锁、令牌、幂等、审计", "SQLite farm.db", "查询与受控仿真写入"),
    ], [1.05, 1.45, 1.55, 1.45, 1.3], font_size=7.9)
    add_h2(doc, "3.2 前端组件关系")
    add_table(doc, ["文件", "职责", "关键约束"], [
        ("spec.js", "产品名称、农场场景、角色导航、指标、数字孪生图层与 Agent 定义", "属于产品合同，不应由页面零散复制"),
        ("engine.js", "本地状态、规则、任务、灌溉、AI 与仿真控制", "输出必须标注 simulated；缺证据返回 NO_GO"),
        ("app.js", "运行时选择、页面路由、交互、数据调用与可访问性", "完整后端失败时不能静默回落并掩盖故障"),
        ("wall.js", "三角色态势大屏与全屏交互", "复用业务状态，不另造数据真相"),
        ("equipment-catalog.js", "农机、厂家、协议与全季装备目录", "兼容性不等于已认证接入"),
        ("mxsj-agents.js", "多 Agent 协同与专家角色扩展", "结论仍需人工门禁"),
        ("styles.css", "布局、密度、主题、响应式和户外可读性", "150% 缩放不得依赖根节点 zoom"),
    ], [1.45, 3.15, 2.25], font_size=8.0)
    add_h2(doc, "3.3 目标云边端架构")
    add_table(doc, ["层级", "目标能力", "候选技术或协议", "当前状态"], [
        ("感知层", "土壤、气象、作物、无人机、卫星、农机遥测", "MQTT、Modbus、CAN、ISOBUS、RTK", "界面与数据合同已预留"),
        ("边缘层", "网关、边缘推理、断网缓存、协议转换、现场安全联锁", "MQTT Gateway、Edge AI、Store-and-forward", "未交付"),
        ("数据层", "事务、空间、时序、向量、图谱、对象存储", "PostgreSQL、PostGIS、TimescaleDB、Milvus、Neo4j、MinIO", "当前为 SQLite"),
        ("AI 层", "农业大模型、RAG、视觉、预测、仿真", "模型注册、版本、特征与评测体系", "规则与公式占位"),
        ("Agent 层", "农场总控、作物、水肥、视觉、机器人、产量、财务", "事件总线、工具权限、记忆与审计", "前端协作沙箱"),
        ("执行层", "水肥一体化、无人机、无人农机、机器人", "签名命令、双人复核、确认回执、物理急停", "禁止真实执行"),
        ("业务层", "SaaS、数据资产、Agent 与模型市场", "租户、计费、权限、合规与 SLA", "产品规划"),
    ], [0.75, 2.2, 2.4, 1.45], font_size=7.8)
    add_h3(doc, "目标消息主题")
    add_code(doc, "farm/{farmId}/device/{deviceId}/data")
    add_code(doc, "farm/{farmId}/device/{deviceId}/cmd")
    add_code(doc, "farm/{farmId}/agent/event")

    # 4
    add_h1(doc, "4 核心业务闭环")
    add_h2(doc, "4.1 从观察到复盘")
    add_table(doc, ["阶段", "系统动作", "责任人", "必须留下的证据"], [
        ("观察", "汇集田块、作物、土壤、气象、设备和任务状态", "数据与田间人员", "值、单位、时间、来源、质量码；空间数据还需 CRS 与版本"),
        ("研判", "规则或 Agent 形成建议并标明不确定性", "农业专家或模型责任人", "输入快照、规则或模型版本、原因、风险"),
        ("补证", "检查缺失、过期、异常或互相矛盾的数据", "任务负责人", "补测记录、照片、传感器校准、现场确认"),
        ("审核", "验证农艺合理性、装备兼容性和安全边界", "有资质审核者", "身份、资质范围、凭据引用、审核意见"),
        ("派工", "创建仿真任务或生产候选命令", "调度人员", "地块、处方、设备、窗口、限值、幂等键"),
        ("执行与确认", "设备回执、现场观察、异常停止", "设备网关与现场人员", "签名回执、轨迹、告警、停止原因"),
        ("验收与复盘", "核对结果、成本、效果并更新知识", "农场主与专家", "完成量、质量、差异、责任、下一季改进"),
    ], [0.75, 2.35, 1.15, 2.55], font_size=7.9)
    add_h2(doc, "4.2 人工门禁状态")
    add_para(doc, "产品规范要求把建议与命令隔开。推荐的统一状态链如下；当前版本只在沙箱内实现其中一部分，生产系统必须由服务端保存并校验状态转换。")
    add_table(doc, ["状态", "含义", "允许的下一步", "禁止条件"], [
        ("建议待核", "规则或 Agent 已给出候选方案", "补证、驳回", "不能直接下发设备"),
        ("NO_GO", "关键证据不足或安全条件不满足", "补证后重新计算", "不能用人工备注绕过硬门禁"),
        ("待人工审核", "证据齐全但需资质人员确认", "批准、退回、拒绝", "审核者身份或资质未验证"),
        ("仿真排队", "仅进入沙箱任务队列", "运行、停止、复盘", "不得标记为真实执行"),
        ("生产候选", "通过服务端授权，可提交现场网关", "签名下发、撤销", "无双人复核、无地理围栏、无设备安全状态"),
        ("已确认", "设备和现场均回传可核验结果", "验收、复盘", "只有请求成功而无设备回执"),
        ("已验收", "产量、质量、资源和异常完成核对", "归档、知识更新", "证据链不完整"),
    ], [1.0, 2.0, 1.65, 2.15], font_size=8.0)
    add_h2(doc, "4.3 幂等与审计")
    add_bullets(doc, [
        "每次写入使用 Idempotency-Key 或服务端生成的请求标识；重复请求不得重复创建任务或重复控制。",
        "审计至少记录农场、操作者、动作、资源、变更前后、原因、是否仿真和时间。",
        "当前 SQLite 审计适合本地追踪；生产环境应写入不可篡改、跨主机、可保留的审计设施。",
    ])

    # 5
    add_h1(doc, "5 功能模块")
    add_h2(doc, "5.1 页面路由总览")
    add_table(doc, ["模块域", "页面键", "主要功能", "当前数据性质"], [
        ("指挥与田块", "dashboard / twin / tasks", "本周农事、田块态势、图层、任务审核与进度", "固定仿真 + 本地规则"),
        ("精准种植", "plants / water / devices", "植株精细管控、水肥研判、传感与设备状态", "仿真观测；处方缺证据即 NO_GO"),
        ("全季装备", "season / fleet / vendors / robots", "播种到收获装备、机队协同、厂家接入、机器人沙箱", "目录与仿真任务"),
        ("收贮加工", "postharvest / history", "入仓、干燥初加工、批次溯源、历年收成", "固定批次与档案数据"),
        ("智能协作", "collab / diagnosis / workbench / agents / ai", "联合诊断、协作工作台、多 Agent、AI 问答与分析", "规则、协作状态与可选外部模型"),
        ("数据底座", "assets / arch", "数据资产、模型与知识登记、云边端架构", "元数据与目标蓝图"),
    ], [1.0, 1.65, 2.65, 1.5], font_size=8.0)
    add_h2(doc, "5.2 关键能力说明")
    add_table(doc, ["能力", "用户动作", "系统反馈", "完成判定"], [
        ("本周农事", "查看、采纳、改期、跳过、停止或回放", "优先级、责任人、窗口、原因、风险与证据", "状态变化可追溯，危险动作不越过门禁"),
        ("田间一张图", "搜索地块、切换图层、查看传感与设备", "同一坐标空间内的地块和对象详情", "无 CRS 时明确禁止导航、面积结算与处方执行"),
        ("水肥研判", "选择田块并核对土壤、作物、天气和设备", "建议或 NO_GO、缺失证据和风险", "生产处方需真实观测、校准和人工审核"),
        ("多机联作", "选择作业链、机具和窗口", "兼容性、冲突、路线与安全检查", "生产命令需设备能力、围栏、RTK 与现场联锁"),
        ("收贮溯源", "查看批次、仓位、质量和异常", "从田块到仓储的批次链", "关键节点具有时间、人员和检测证据"),
        ("AI 问答", "以自然语言询问政策、农艺或系统状态", "带边界、证据和风险的答复", "不能替代执业判断或自动下发控制"),
        ("数据资产", "检索数据集、知识和模型", "版本、来源、负责人、质量与用途", "生产数据需血缘、权限、保留和质量 SLA"),
    ], [1.05, 2.0, 2.25, 1.5], font_size=7.8)
    add_h2(doc, "5.3 多 Agent 协同")
    add_para(doc, "目标 Agent 角色包括农场总控、作物、水肥、视觉、机器人、产量和财务。协同应由任务编排器分配工具和数据范围，所有输出带来源、模型或规则版本、置信度和下一步建议。高风险结论由农艺、农机、安全和业务负责人联合审核。")
    add_bullets(doc, [
        "Agent 只能调用其声明的工具和农场范围，不得自行扩大权限。",
        "冲突意见不做简单多数表决；先比较证据质量、专业资质和风险等级。",
        "记忆内容必须区分事实、推断、用户输入和模型生成，过期事实自动降级。",
        "涉及药剂、剂量、设备运动或资金结算时，必须转入人工门禁。",
    ])

    # 6
    add_h1(doc, "6 首页与态势大屏")
    add_h2(doc, "6.1 首页设计")
    add_para(doc, "首页以“先处理任务，再理解背景”为信息顺序。顶部工作区消除无意义留白，菜单与内容起点对齐；左侧导航保持分组清楚，中部先显示关键任务，地图和态势作为决策背景，右侧动作区保持可达。")
    add_table(doc, ["区域", "内容", "交互要求", "风险控制"], [
        ("全局顶栏", "农场、搜索、运行模式、天气、时钟、刷新、主题、大屏、通知、角色", "键盘可达、状态可读、窄宽度下保持核心动作", "运行模式始终可见，避免把本地仿真误当在线生产"),
        ("任务队列", "今天和本周待办、优先级、责任人、截止时间", "先展示高优，支持渐进展开和状态动作", "动作前展示证据与影响；不可只靠颜色传达等级"),
        ("田块地图", "地块、传感、设备、风险和作物图层", "搜索、图层、详情与键盘操作", "坐标基准缺失时禁止路线与面积结论"),
        ("态势摘要", "水肥、病虫、设备、作业、仓储与异常", "从摘要跳转到责任页面", "指标标明目标、仿真或真实来源"),
    ], [1.0, 2.15, 2.0, 1.65], font_size=7.8)
    add_h2(doc, "6.2 三类态势大屏")
    add_table(doc, ["大屏", "首屏问题", "优先指标", "下钻方向"], [
        ("农场主态势", "今天有什么必须处理，是否影响产量与成本", "任务、天气窗口、水肥、设备、预计进度", "任务、田块、设备、收贮"),
        ("专家态势", "哪里需要研判，哪些证据不足", "风险地块、缺证据项、审核队列、模型分歧", "联合诊断、水肥、植株、协作工作台"),
        ("监管态势", "哪些异常需要核查，效果证据是否完整", "高优异常、节水目标、全季进度、批次溯源、设备合规", "取证、档案、数据资产与系统状态"),
    ], [1.0, 2.25, 2.1, 1.45], font_size=8.0)
    add_h2(doc, "6.3 视觉与可访问性")
    add_bullets(doc, [
        "项目包含 12 张本地真实感农业图片，覆盖棉花、玉米、小麦、灌溉、无人机、气象站和土壤探头；这些图片是静态资料，不代表实时视频。",
        "户外可读性以高对比文本、明确边界和适度密度实现；150% 显示缩放通过布局重排适配，不使用根节点整体缩放。",
        "所有图标按钮具有可访问名称；对话框具备正确角色、标题关联、焦点约束和关闭后恢复。",
        "动态刷新不得在用户填写表单或操作对话框时重建界面；支持减少动画偏好。",
    ])

    # 7
    add_h1(doc, "7 数据架构与质量合同")
    add_h2(doc, "7.1 数据分层")
    add_table(doc, ["层级", "内容", "当前载体", "生产要求"], [
        ("主数据", "农场、地块、作物、设备、人员、厂家、农资", "spec.js / SQLite", "统一编码、租户隔离、生命周期与责任人"),
        ("观测数据", "土壤、气象、作物、设备与影像观测", "固定仿真状态 / sensor_readings", "时序库、设备身份、质量码、校准与缺测规则"),
        ("业务事务", "任务、审核、水肥计划、机器人任务、仓储批次", "SQLite 与本地状态", "强一致状态机、幂等、权限与不可变审计"),
        ("知识与模型", "农艺知识、政策、规则、模型和评测", "knowledge / models", "来源、许可、版本、适用区域、评测与回滚"),
        ("空间数据", "地块边界、路线、设备位置、禁入区", "屏幕坐标与 GeoJSON 字段", "PostGIS、CRS、测绘来源、精度、版本与拓扑校验"),
    ], [0.95, 2.15, 1.55, 2.15], font_size=7.9)
    add_h2(doc, "7.2 最小数据合同")
    add_para(doc, "任何进入决策链的观测值至少包含以下字段；缺少必要字段时系统只能显示或请求补证，不能生成生产命令。")
    add_table(doc, ["字段", "说明", "校验规则"], [
        ("value", "原始或处理后的数值", "明确类型、范围与空值语义"),
        ("unit", "单位", "采用统一单位字典，换算有记录"),
        ("observed_at", "观测时间", "使用带时区时间；超过业务时效阈值即降级"),
        ("source", "设备、人工、遥感或模型来源", "可追到设备或数据集版本"),
        ("quality_code", "正常、估算、缺测、异常、校准中等", "低质量值不能静默参与处方"),
        ("farm_id / land_id", "租户与空间范围", "服务端强制校验，不信任前端参数"),
        ("geometry metadata", "CRS、版本、测绘来源和精度", "缺失时禁止导航、面积结算和处方执行"),
    ], [1.35, 2.55, 2.9], font_size=8.1)
    add_h2(doc, "7.3 当前 SQLite 表")
    add_table(doc, ["领域", "数据表", "用途"], [
        ("农场主数据", "farms, lands, crops", "农场、地块、土壤健康、作物品种与生育期"),
        ("任务", "farm_tasks", "任务、地块、负责人、计划时间、优先级与状态"),
        ("设备与遥测", "devices, sensor_readings", "设备登记、位置、状态、协议和观测值"),
        ("智能协同", "agents, agent_tasks", "Agent 名册、工具、评分、目标和结果"),
        ("水肥", "irrigation_plans", "田块、水量、肥料、原因、门禁与计算状态"),
        ("机器人", "robot_missions", "机器人仿真任务与执行状态"),
        ("资产", "data_assets, knowledge, models", "数据集、知识和模型元数据"),
        ("系统", "app_state, audit_events", "当前状态、请求幂等和变更审计"),
    ], [1.05, 2.2, 3.55], font_size=8.2)
    add_h2(doc, "7.4 数据保留与备份建议")
    add_bullets(doc, [
        "本地开发：变更前备份 data/farm.db，数据库文件不与来源不明的仓库内容互相覆盖。",
        "生产遥测：热数据、历史汇总、对象影像和审计分别定义保留期；所有删除遵循审批与可恢复策略。",
        "模型与知识：保留输入快照、版本、评测、批准人和回滚点，确保结论可重现。",
    ])

    # 8
    add_h1(doc, "8 接口与集成")
    add_h2(doc, "8.1 API 设计约定")
    add_bullets(doc, [
        "读取接口返回 data_mode、source、as_of、simulated 等元数据，调用方不得丢弃。",
        "写接口默认返回 423 锁定；仅在显式启用本地仿真写入并提供正确令牌时开放。",
        "请求体上限为 64 KiB；start-web.ps1 的 AI 问答代理限制为 16 KiB，并执行本机速率限制。",
        "写操作使用幂等标识并记录审计；未知设备、非法状态转换和越界输入均失败关闭。",
        "CORS、Host、Origin、路径和内容类型执行白名单检查；系统只建议绑定 127.0.0.1。",
    ])
    add_h2(doc, "8.2 设备与边缘接入")
    add_table(doc, ["接入对象", "协议或数据", "接入前验证", "失效策略"], [
        ("土壤与气象传感器", "MQTT / Modbus；含值、单位、时间、质量码", "设备身份、校准记录、采样频率、时间同步", "标记缺测或异常，不自动用旧值替代"),
        ("拖拉机与农具", "ISOBUS / CAN / 厂家 API", "能力描述、机具匹配、软件版本、安全功能", "停止下发，保留现场人工控制"),
        ("无人机", "厂家任务 API、航线、影像和遥测", "空域、飞手、气象、电量、返航和禁飞区", "返航或安全降落由设备侧保障"),
        ("泵站与水肥机", "PLC / Modbus / 网关命令", "阀门映射、压力流量、最大剂量、现场联锁", "通信丢失时进入设备定义的安全状态"),
        ("机器人", "任务、轨迹、障碍和状态回执", "地理围栏、定位质量、避障、自检和现场急停", "超时即停止并请求人工介入"),
    ], [1.15, 1.65, 2.45, 1.55], font_size=7.7)
    add_h2(doc, "8.3 厂家接入验收包")
    add_numbered(doc, [
        "设备型号、固件、能力、命令字典、量程、单位和错误码。",
        "身份与密钥生命周期、传输加密、签名或消息认证、重放防护。",
        "正常、断网、延迟、重复、乱序、异常值、断电和恢复测试记录。",
        "安全停机、现场手动接管、物理急停和责任边界说明。",
        "仿真、台架、小范围田测、扩大试运行和生产批准的逐级证据。",
    ])

    # 9
    add_h1(doc, "9 安全与权限")
    add_h2(doc, "9.1 当前控制")
    add_table(doc, ["控制点", "当前实现", "有效范围", "生产差距"], [
        ("写入默认拒绝", "AI_FARM_ALLOW_DEMO_WRITES 未开启时返回 423", "本地 FastAPI 沙箱", "需要策略引擎、审批和按资源授权"),
        ("仿真写入令牌", "X-AI-Farm-Token 与环境变量常量时间比较", "本机联调", "不是用户身份、会话或多租户授权"),
        ("输入限制", "64 KiB API 上限；AI 代理 16 KiB；Pydantic 校验", "HTTP 请求", "还需文件扫描、字段级规则和全链路配额"),
        ("浏览器边界", "CSP、禁止嵌入、nosniff、引用与权限策略", "本地页面", "生产需 HTTPS、严格 CSP、供应链完整性"),
        ("来源白名单", "CORS、Host、Origin 与本地回环地址限制", "本机服务", "生产需网关、WAF、服务身份与零信任网络"),
        ("幂等与审计", "请求标识和 audit_events", "SQLite 内部", "需不可变、异地保留、告警与审计查询权限"),
    ], [1.25, 2.3, 1.35, 1.9], font_size=7.6)
    add_h2(doc, "9.2 权限模型")
    add_para(doc, "前端角色选择器只用于展示不同工作台，不能作为授权依据。生产系统应组合 RBAC 与 ABAC：RBAC 定义岗位，ABAC 再校验租户、农场、地块、设备、作业类型、时间窗、资质和风险等级。")
    add_table(doc, ["资源", "读取", "建议或审核", "生产执行"], [
        ("农场与田块", "按租户与农场范围", "专家需被指派且资质适用", "不适用"),
        ("水肥处方", "农场主、专家、监管按范围", "农艺专家审核并保留证据", "双人复核 + 设备网关策略"),
        ("农机与机器人", "状态按范围可见", "农机专家确认兼容和路线", "调度授权 + 现场安全条件 + 设备签名"),
        ("数据资产与模型", "按敏感级别和用途", "数据或模型负责人审批", "发布需评测、签名和回滚"),
        ("审计", "监管和审计角色只读", "不能修改历史", "独立保留与导出审批"),
    ], [1.25, 1.65, 2.0, 1.9], font_size=7.8)
    add_h2(doc, "9.3 重点威胁")
    add_bullets(doc, [
        "伪造或重放控制命令：采用短时签名、单次 nonce、设备身份、序列号和幂等检查。",
        "跨农场越权：所有查询和写入由服务端从身份上下文派生 farm_id，不接受前端自报作为唯一依据。",
        "不可信传感器与数据投毒：设备证书、校准、异常检测、来源评分和人工复核共同控制。",
        "模型提示注入或知识污染：工具最小权限、检索来源白名单、内容隔离、输出策略和审计。",
        "供应链与依赖风险：锁定版本、生成物料清单、扫描依赖与构建产物、签名发布。",
        "本地密钥泄漏：.env 不提交仓库、不打印密钥；生产使用密钥管理服务并定期轮换。",
    ])

    # 10
    add_h1(doc, "10 农业决策与农机安全")
    add_h2(doc, "10.1 农艺联合审核")
    add_para(doc, "农业建议不能只看单一传感值。系统应把生育期、土壤、天气、作物长势、病虫、作业窗口、设备能力和历史效果放入同一证据包，再由相应专家审核。")
    add_table(doc, ["作业", "必需输入", "NO_GO 条件", "验收证据"], [
        ("灌溉施肥", "土壤含水率、根区深度、作物阶段、天气与有效降雨、流量压力、肥液兼容", "观测过期或质量低、无设备能力、剂量超限、人员未审核", "流量、时长、压力、EC/pH、地块反馈与异常"),
        ("植保施药", "病虫证据、药剂标签、作物与生育期、风雨、缓冲区、喷头与校准", "无确诊证据、超标签剂量、风速或天气不合格、禁入区冲突", "药剂批次、剂量、航迹或轨迹、天气、覆盖与复查"),
        ("播种", "品种、播期、土壤温湿度、密度、深度、机具能力和 RTK", "土壤条件不合格、种子或机具不匹配、定位质量不足", "速度、深度、株距、漏播重播、面积与种子用量"),
        ("收获", "成熟度、水分、天气、损失阈值、机具和仓储能力", "水分或天气不合格、仓容不足、设备故障", "收获量、损失率、水分、杂质、批次与入仓"),
        ("仓储初加工", "批次、含水率、温湿度、仓位、设备能力和质量标准", "批次不清、质量超限、仓位或设备不合格", "入出仓、温湿度曲线、干燥参数、质量检验与异常"),
    ], [0.9, 2.45, 2.0, 1.45], font_size=7.25)
    add_h2(doc, "10.2 水肥计算边界")
    add_para(doc, "生产灌溉可采用 ETc = Kc × ET0 作为需水计算的一部分，再结合有效降雨、根区可利用水、目标含水区间和系统效率形成净灌溉量。该公式必须使用当地标定的 Kc、气象、土壤和设备参数。当前仓库只有仿真观测，因此后端始终不从这些数据计算可执行水量，而是返回 NO_GO。")
    add_h2(doc, "10.3 农机功能安全")
    add_numbered(doc, [
        "处方地块必须具有可验证的 CRS、边界版本、精度与禁入区；屏幕坐标不能转为导航路线。",
        "核对动力机与农具的接口、功率、液压、PTO、通信、幅宽、速度和安全功能，目录匹配不等于认证。",
        "自动作业前检查定位质量、地理围栏、障碍物、人车隔离、天气、油电、制动、转向和通信。",
        "命令含农场、地块、设备、处方版本、有效期、限值、操作者、审批链、序列号和签名。",
        "设备侧必须独立实施超速、越界、失联、障碍和急停策略；云端远程停止不能替代物理急停。",
        "执行后接收设备签名回执与现场确认；无回执、回执过期或轨迹异常时不得自动验收。",
    ])
    add_h2(doc, "10.4 专家联审角色")
    add_table(doc, ["专家", "主要判定", "一票否决情形"], [
        ("农艺与作物", "作物阶段、品种、营养、病虫与预期效果", "违反药剂标签、证据不足或明显伤害作物"),
        ("水利与土壤", "水量、肥量、盐分、根区和水力能力", "来源不可靠、剂量越限或系统能力不足"),
        ("农机与安全", "机具兼容、路线、定位、围栏、联锁与急停", "人员或障碍风险、定位不合格、无现场接管"),
        ("数据与软件", "数据合同、版本、权限、幂等、审计和可回滚", "来源不可追、跨租户、无权限或状态机被绕过"),
        ("测试与可靠性", "正常与失效场景、回归、恢复和证据完整", "P0 测试失败或无法证明失败关闭"),
    ], [1.2, 3.25, 2.35], font_size=8.0)

    # 11
    add_h1(doc, "11 部署运行")
    add_h2(doc, "11.1 推荐本地启动")
    add_para(doc, "本地评审优先使用项目自带 PowerShell 启动器。它只绑定回环地址并服务前端，可选代理 AI 问答。")
    add_code(doc, "cd \"C:\\Users\\dongr\\Desktop\\新建文件夹\\ai-farm-os\"")
    add_code(doc, "powershell -ExecutionPolicy Bypass -File .\\start-web.ps1")
    add_code(doc, "浏览器访问 http://127.0.0.1:8080/")
    add_h2(doc, "11.2 其他运行方式")
    add_table(doc, ["方式", "入口", "适用场景", "限制"], [
        ("直接打开", "打开系统.html", "快速查看界面", "没有本地 API 与 AI 代理能力"),
        ("本地静态服务", "start-web.ps1", "推荐的界面评审与可选 AI 问答", "业务数据仍以 FarmEngine 为主"),
        ("后端沙箱", "start.bat 或 uvicorn backend.app:app", "接口联调、SQLite 与安全门禁测试", "business_api=partial，写入默认锁定"),
    ], [1.1, 1.7, 2.35, 1.65], font_size=8.1)
    add_h2(doc, "11.3 可选 AI 配置")
    add_para(doc, "在项目根目录 .env 中配置以下变量后，由 start-web.ps1 在本机代理问答。密钥不得写入前端、文档或版本库。")
    add_code(doc, "OPENAI_API_KEY=<secret>")
    add_code(doc, "OPENAI_MODEL=<approved-model>")
    add_code(doc, "OPENAI_BASE_URL=<approved-endpoint>")
    add_h2(doc, "11.4 受控仿真写入")
    add_para(doc, "仅用于本机接口测试。令牌不是生产身份系统，完成测试后应关闭相关进程并移除环境变量。")
    add_code(doc, "$env:AI_FARM_ALLOW_DEMO_WRITES='1'")
    add_code(doc, "$env:AI_FARM_WRITE_TOKEN='<random-local-token>'")
    add_code(doc, "python -m uvicorn backend.app:app --host 127.0.0.1 --port 8090")
    add_h2(doc, "11.5 运行检查")
    add_bullets(doc, [
        "地址栏必须是 127.0.0.1 或 localhost；不得把开发服务直接暴露到局域网或互联网。",
        "顶栏运行模式应显示本地仿真或后端沙箱，不能出现暗示生产在线的文案。",
        "检查浏览器控制台、网络请求、健康接口和页面主要动作；缓存导致旧版时执行强制刷新。",
        "生产部署必须另行提供 HTTPS、反向代理、身份系统、密钥管理、日志、监控、备份和灾难恢复。",
    ])

    # 12
    add_h1(doc, "12 测试与验收")
    add_h2(doc, "12.1 发布门禁")
    add_para(doc, "项目统一验证入口为 verify.ps1。每次修改首页、态势大屏、业务规则、后端安全或启动脚本后，都应完整执行，而不是只运行单个测试文件。")
    add_code(doc, "powershell -NoProfile -ExecutionPolicy Bypass -File .\\verify.ps1")
    add_table(doc, ["门禁", "覆盖内容", "2026-09-14 结果"], [
        ("JavaScript 语法", "frontend/assets 下脚本", "通过"),
        ("Python 编译", "backend 与 tests", "通过"),
        ("PowerShell 解析", "start-web.ps1 与 verify.ps1", "通过"),
        ("Pytest", "安全、农业门禁、前端合同、可访问性与界面约束", "34 项通过"),
        ("第三方告警", "Starlette TestClient 使用的 anyio 别名弃用提示", "1 条；不影响本次结果，后续升级依赖时处理"),
    ], [1.35, 3.55, 1.9], font_size=8.2)
    add_h2(doc, "12.2 自动化测试重点")
    add_bullets(doc, [
        "安全：默认锁定写入、令牌校验、请求体上限、CORS 拒绝、未知设备与非法状态失败关闭。",
        "农业：灌溉证据不足时稳定返回 NO_GO，仿真观测不被改写，也不计算可执行剂量。",
        "事实标记：健康接口、页面与导出明确 partial、simulated、目标指标和坐标限制。",
        "前端：版本化资源、语义导航、可访问名称、对话框焦点、大屏可操作、任务优先首页、地图坐标统一。",
        "可靠性：刷新不破坏表单、错误不静默回落、内容安全输出、150% 缩放与窄宽度重排。",
        "视觉素材：使用本地真实感农业媒体，避免占位素材或把资料图冒充实时画面。",
    ])
    add_h2(doc, "12.3 人工验收清单")
    add_table(doc, ["验收项", "操作", "通过条件"], [
        ("启动", "运行 start-web.ps1 并访问首页", "页面无阻塞错误；模式标识清楚"),
        ("首页", "检查顶栏、菜单、任务、地图和动作区", "菜单上方无空白；任务先于背景信息；布局对齐"),
        ("角色", "切换农场主、专家、监管", "导航、用语和重点变化正确；不暗示获得真实权限"),
        ("地图", "搜索地块、切图层、查看对象", "对象在同一屏幕坐标空间；缺 CRS 警示可见"),
        ("水肥", "进入田块研判并尝试形成方案", "缺证据时 NO_GO；不给可执行水量"),
        ("设备与机器人", "登记或派发仿真任务", "显示 simulated / sandbox；不声称真实设备执行"),
        ("大屏", "打开、键盘操作并退出", "三角色内容正确；焦点受控；退出恢复原位置"),
        ("显示缩放", "Windows 或浏览器设为 150%", "核心任务和动作仍可见，无整体缩小或横向溢出"),
        ("故障", "断开接口或输入非法数据", "明确报错并停止，不静默使用旧数据伪装成功"),
        ("审计", "使用幂等键重复提交本地仿真写入", "只产生一次业务变更，审计记录可追踪"),
    ], [1.0, 2.45, 3.35], font_size=7.9)
    add_h2(doc, "12.4 缺陷分级")
    add_table(doc, ["级别", "定义", "发布规则"], [
        ("P0", "越权、真实设备误控制、剂量或路线错误、数据泄露、无法安全停止", "立即停止发布，修复并完成安全回归"),
        ("P1", "核心任务无法完成、结论不一致、审计断链、主要角色或大屏不可用", "发布前修复；需专项回归"),
        ("P2", "次要交互、文案、样式或低频兼容问题", "评估影响后排期，不得掩盖事实边界"),
    ], [0.7, 4.55, 1.55], font_size=8.1)

    # 13
    add_h1(doc, "13 生产化路线")
    add_para(doc, "生产化不是把当前沙箱接上设备即可。必须先补齐可信身份、真实数据、空间基准、设备安全和可追溯验收，再逐步扩大作业范围。")
    add_table(doc, ["阶段", "优先交付", "退出条件"], [
        ("P0 信任与安全", "统一身份、多租户与农场隔离、服务端授权、设备证书与签名、命令网关、双人复核、物理联锁、不可变审计、密钥管理", "安全评审通过；故障注入证明默认拒绝；可从云端到设备回执完整追踪"),
        ("P1 数据与农艺", "真实遥测、PostGIS/时序存储、CRS 与测绘、数据质量、处方模型、本地 Kc/土壤/设备标定、影像评测", "试验设计完成；模型达到预先定义的安全和效果阈值；专家签署适用范围"),
        ("P2 可靠运行", "边缘缓存、断网策略、观测性、告警、备份恢复、容量与高可用、设备数字孪生", "SLA、RTO/RPO、恢复演练和季节性负载测试通过"),
        ("P3 规模优化", "跨农场调度、成本与产量优化、模型与 Agent 生命周期、供应商生态和合规报表", "多季、多区域证据证明可复制，且风险与收益可量化"),
    ], [1.0, 3.35, 2.45], font_size=7.7)
    add_h2(doc, "13.1 生产首批试运行建议")
    add_numbered(doc, [
        "先只读：接入真实传感、设备状态和空间数据，验证质量、时效、坐标与告警。",
        "再建议：让系统生成处方候选，但由现有人工流程执行，并比较建议与专家结果。",
        "再台架：在隔离设备和小范围台架验证签名、限值、重复请求、断网和急停。",
        "再小区：选择低风险地块、白天窗口和现场监护，逐项收集执行与验收证据。",
        "再扩大：只有连续多个窗口满足安全、农艺和可靠性阈值，才扩大设备、面积与自动化等级。",
    ])
    add_h2(doc, "13.2 目标指标治理")
    add_bullets(doc, [
        "节水、识别准确率、巡田减少和亩均增收都应定义基线、样本、时间范围、计算公式、置信区间和责任人。",
        "界面同时展示目标、当前观测和数据质量，不能把目标值绘制成已实现结果。",
        "每季冻结指标口径并保存原始数据；任何模型或口径变更都重新计算并保留差异。",
    ])

    # 14
    add_h1(doc, "14 运维与故障处理")
    add_h2(doc, "14.1 例行检查")
    add_table(doc, ["频率", "检查内容", "异常动作"], [
        ("每次启动", "健康状态、运行模式、控制台、关键页面、时间与天气来源", "停止使用错误页面，保留日志并回退到最近可验证版本"),
        ("每日", "高优任务、缺测、过期观测、设备离线、审计写入、备份", "创建责任明确的故障任务；高风险业务进入 NO_GO"),
        ("每周", "规则与模型变更、依赖告警、权限、容量、失败请求、恢复抽查", "完成回归并关闭过期账号或密钥"),
        ("作业前", "天气、地块、处方、机具、自检、围栏、通信、现场人员与急停", "任何硬条件不满足即取消自动作业"),
        ("每季", "指标口径、模型标定、设备校准、知识版本、灾备与安全演练", "更新适用范围和风险清单后再开季"),
    ], [0.9, 3.45, 2.45], font_size=7.9)
    add_h2(doc, "14.2 常见故障")
    add_table(doc, ["现象", "可能原因", "处理方法"], [
        ("界面仍是旧版", "浏览器缓存或旧服务未退出", "确认端口进程，停止旧服务，重新启动并强制刷新"),
        ("写入返回 423", "仿真写入未显式开启", "这是安全默认；仅在本机联调时设置开关和随机令牌"),
        ("写入返回 401", "缺少或错误的 X-AI-Farm-Token", "核对本机环境变量和请求头，不记录令牌值"),
        ("AI 问答不可用", "未配置密钥、模型或端点，或上游网络故障", "查看 /api/ai/status 与本地日志；保持规则引擎可用并显示明确错误"),
        ("业务仍使用本地引擎", "后端报告 business_api=partial", "属于当前预期；补全后端并通过完整合同测试后才可改为 complete"),
        ("地图不能生成路线", "只有屏幕坐标或 CRS、精度、版本缺失", "补充经过验证的空间数据；不得手工绕过"),
        ("任务重复", "调用方重试未复用幂等键", "对同一业务意图复用 Idempotency-Key 并检查审计"),
    ], [1.45, 2.05, 3.3], font_size=8.0)
    add_h2(doc, "14.3 变更管理")
    add_bullets(doc, [
        "所有功能变更说明影响角色、页面、接口、数据合同、安全门禁和回滚方式。",
        "农业规则变更由相应专家评审；农机控制变更增加台架和失效测试；安全边界变更需独立复核。",
        "发布包保存版本、提交、依赖、测试结果和已知限制；发生问题时按证据回滚，不覆盖用户数据。",
    ])

    # Appendix A
    add_h1(doc, "附录 A 接口清单", new_page=True)
    api_rows = [
        ("GET", "/api/health", "运行状态、规范版本、数据模式、业务 API 完整度与控制模式", "只读"),
        ("GET / POST", "/api/role", "读取或设置界面角色", "POST 为仿真写入；不构成授权"),
        ("GET", "/api/dashboard", "首页任务、态势和目标指标", "只读"),
        ("GET", "/api/architecture", "当前与目标架构信息", "只读"),
        ("GET", "/api/twin", "田块、设备、图层与空间元数据", "只读；无 CRS 禁止执行"),
        ("GET", "/api/farms", "农场列表", "只读"),
        ("GET", "/api/lands", "地块列表与状态", "只读"),
        ("GET", "/api/agents", "Agent 名册与状态", "只读"),
        ("POST", "/api/agents/dispatch", "创建 Agent 仿真任务", "写入锁 + 令牌 + 审计"),
        ("POST", "/api/agents/orchestrate", "多 Agent 编排", "仿真；需人工审核"),
        ("GET", "/api/devices", "设备和传感状态", "只读"),
        ("POST", "/api/devices/register", "登记沙箱设备", "不建立真实设备连接"),
        ("POST", "/api/devices/{device_id}/control", "创建设备仿真控制请求", "executable=false"),
        ("GET", "/api/robots", "机器人与任务状态", "只读"),
        ("POST", "/api/robots/dispatch", "机器人任务排队", "等待人工确认；无真实命令"),
        ("GET", "/api/tasks", "按角色读取任务", "只读"),
        ("POST", "/api/tasks", "创建仿真任务草稿", "写入锁 + 令牌 + 审计"),
        ("POST", "/api/tasks/{task_id}/status", "受控任务状态转换", "校验状态机与幂等"),
        ("POST", "/api/tasks/{task_id}/audit", "提交专家审核", "身份、资质与证据声明"),
        ("GET", "/api/irrigation", "读取水肥研判", "当前数据不产生可执行剂量"),
        ("POST", "/api/irrigation/apply", "提交水肥方案候选", "证据不足稳定 NO_GO"),
        ("POST", "/api/ai/chat", "农业与系统问答", "本地规则或可选外部模型"),
        ("POST", "/api/ai/vision", "关键词规则视觉判断", "不是图像诊断模型"),
        ("GET", "/api/ai/yield", "未标定的产量公式", "仅供沙箱分析"),
        ("GET", "/api/ai/rag", "知识检索结果", "保留来源与适用范围"),
        ("GET", "/api/ai/risk", "风险仿真", "不作为自动执行依据"),
        ("GET", "/api/assets", "数据资产目录", "只读"),
        ("GET", "/api/models", "模型目录", "只读"),
        ("GET", "/api/knowledge", "知识目录", "只读"),
    ]
    add_table(doc, ["方法", "路径", "用途", "边界"], api_rows, [0.8, 2.4, 2.45, 1.15], font_size=7.1)

    # Appendix B
    add_h1(doc, "附录 B 数据表与代码索引")
    add_h2(doc, "B.1 数据表索引")
    add_table(doc, ["表", "关键内容", "数据责任"], [
        ("farms", "名称、区域、面积、作物重点、边缘节点与 MQTT 状态", "农场主数据管理员"),
        ("lands", "地块编码、面积、土壤、健康、水分、养分、风险、GeoJSON", "GIS 与农艺负责人"),
        ("crops", "作物、品种、生育期、播期与预期产量", "农艺负责人"),
        ("farm_tasks", "任务、地块、负责人、状态、时间与优先级", "生产调度"),
        ("devices", "设备、类型、状态、位置、协议、电量与厂家", "设备管理员"),
        ("sensor_readings", "指标、值、单位与时间", "数据平台与设备管理员"),
        ("agents / agent_tasks", "Agent、工具、目标、状态与结果", "AI 平台负责人"),
        ("irrigation_plans", "水肥候选、原因、状态与计算门禁", "水肥专家"),
        ("robot_missions", "机器人任务和状态", "机器人调度"),
        ("data_assets / knowledge / models", "数据、知识和模型元数据", "数据与模型负责人"),
        ("app_state", "本地应用状态", "系统维护者"),
        ("audit_events", "请求、操作者、动作、资源、前后值、原因和仿真标记", "安全与审计负责人"),
    ], [2.05, 3.35, 1.4], font_size=7.9)
    add_h2(doc, "B.2 代码与规范索引")
    add_table(doc, ["路径", "用途"], [
        ("README.md", "项目定位、启动方式、运行模式和安全声明"),
        ("docs/FINAL_SPEC.md", "权威产品终版规范、目标架构、模块和验收要求"),
        ("docs/EXPERT_COUNCIL.md", "多专家职责、联审流程和发布门禁"),
        ("frontend/index.html", "单页应用结构、顶栏、导航和对话框容器"),
        ("frontend/assets/spec.js", "产品、角色、导航、目标指标、图层与 Agent 合同"),
        ("frontend/assets/engine.js", "本地规则、仿真数据、任务和控制沙箱"),
        ("frontend/assets/app.js", "运行时、页面路由、交互和数据调用"),
        ("frontend/assets/wall.js", "三角色态势大屏"),
        ("frontend/assets/styles.css", "视觉系统、布局、密度、主题和可访问性"),
        ("backend/app.py", "FastAPI 接口、安全中间件、写入门禁与审计"),
        ("backend/db.py", "SQLite 表结构、连接和场景初始化"),
        ("start-web.ps1", "本地静态服务、AI 代理与 Host/路径保护"),
        ("verify.ps1", "统一发布验证入口"),
        ("tests/", "后端安全、农业门禁、前端合同和界面约束测试"),
    ], [2.55, 4.25], font_size=8.0)
    add_h2(doc, "B.3 术语")
    add_table(doc, ["术语", "定义"], [
        ("NO_GO", "证据、权限或安全条件不满足，系统拒绝形成可执行结果"),
        ("仿真", "只改变沙箱状态，不代表真实设备、田块或仓储已经发生变化"),
        ("数据质量码", "描述观测是否正常、估算、缺测、异常或处于校准中的字段"),
        ("CRS", "坐标参考系统；决定空间数据能否正确叠加、测量和导航"),
        ("幂等", "同一业务请求重复发送时只产生一次效果"),
        ("物理急停", "设备现场独立安全回路；不能被云端按钮替代"),
        ("受控作业沙箱", "用于验证流程、规则和接口，但禁止把输出视为真实生产执行"),
    ], [1.55, 5.25], font_size=8.2)
    add_para(doc, "文档结束。后续若代码、产品规范、接口完整度或农业安全规则发生变化，应同步更新本说明书并重新执行完整验证。", italic=True)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(str(OUTPUT))


if __name__ == "__main__":
    build_document()
